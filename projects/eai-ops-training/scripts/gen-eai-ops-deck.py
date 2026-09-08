# -*- coding: utf-8 -*-
"""小象智汇 · 成绩执行器 运营培训教材 deck 生成脚本（可复跑）。

== Step 1 受众三问 ==
给谁看：代理商运营人员 / 装机负责人（外行，不写代码，可能没用过平台）。
观众已知：自己机构的 e听说 账号密码；学校/班级等业务信息。
看完能复述三句话：
  1. 网页 eai.xiaoxiangzhihui.com 是唯一入口：配 e听说 账密、发作业、收成品。
  2. 装机只做一次：下载安装包，一路下一步，向导里什么都不用填。
  3. 之后全自动：改密码只改网页凭据卡，装机电脑零操作。

== 叙事大纲（主线：8 分钟完成装机配置，之后日常导出全自动）==
P1  封面（主张句）           → 立论：装一次，之后全自动
P2  三答卡片                 → 是什么 / 你要做什么 / 装完之后
P3  全景链路                 → 一单导出怎么流转，你只碰网页
P4  章节 01 网页配置 3 分钟   → 承上：先做网页侧
P5  登录（截图圈注）         → 步骤 1
P6  进装机页（截图圈注）     → 步骤 2
P7  凭据卡+下载（截图圈注）  → 步骤 3-4
P8  章节 02 安装 3 分钟
P9  向导示意页               → 三项全不填，点安装
P10 装完即自动运行（4 卡）    → 打消「以后要不要管它」
P11 章节 03 验证 2 分钟
P12 日志判定（示意面板）      → 绿了收工 / 红了找平台
P13 章节 04 日常使用 全自动
P14 发起导出（截图圈注）     → 日常唯一动作
P15 收成品（截图圈注）       → 闭环：网页收文件
P16 两个日常场景对比         → 改密码 / 新增代理商
P17 尾页三句话               → 复述收束

== 素材出处 ==
截图 2026-09-08 实拍 https://eai.xiaoxiangzhihui.com（zhang.ops 演示登录，账号已脱敏）；
流程口径同项目 DEPLOYMENT.md §6/§7 与网页「执行器下载」页自述；安装包 v0.3.0。
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import primitives as P
from pptx.util import Inches, Emu

P.use_theme("warm")  # 暖珊瑚主题（用户指定暖色）

HERE = os.path.dirname(os.path.abspath(__file__))
PROJ = os.path.dirname(HERE)
SHOTS = os.path.join(PROJ, "assets", "screenshots")
OUT = os.path.join(PROJ, "deck-eai-ops-training.pptx")

META = {
    "kicker": "OPERATIONS TRAINING · 运营培训教材",
    "title": "装一次，之后全自动",
    "subtitle": "小象智汇 · 成绩执行器 装机与使用培训",
    "tagline": "面向代理商运营 / 装机负责人 · 装机全程约 8 分钟 · 不需要懂技术",
    "domain": "eai.xiaoxiangzhihui.com",
    "version": "安装包 v0.3.0",
    "date": "2026-09-08",
    "author": "小象智汇平台方",
}


def media(idx, label, title, img, sidebar, caption=None, footnote=None):
    return P.slide_media(P.prs, idx, label, title, os.path.join(SHOTS, img),
                         bullets=None, caption=caption, sidebar=sidebar,
                         footnote=footnote)


# ── P1 封面 ──────────────────────────────────────────────────────────
P.prs = P.new_deck()
P.slide_cover(P.prs, META)
P.notes(P.prs.slides[-1],
        "今天这份教材讲一件事：把成绩执行器装到代理商的电脑上，装完之后日常导出全自动。"
        "装机加配置全程 8 分钟，不需要懂技术。")

# ── P2 开篇三答 ──────────────────────────────────────────────────────
P.slide_cards(P.prs, 2, "Overview", "三件事，一次讲清", [
    ("这套系统是什么", "平台网页 + 自动执行器。\n网页上发单，代理商电脑里的\n执行器自动登 e听说 下载\n成绩、美化、传回平台。"),
    ("你要做什么", "只有两件事：\n网页上配好 e听说 账密，\n电脑上装一次执行器。\n全程点鼠标，约 8 分钟。"),
    ("装完之后", "这台电脑不再需要人碰。\n网页发起导出，几分钟后\n回网页收成品文件。"),
], cols=3)
P.notes(P.prs.slides[-1],
        "先定位：它是网页加自动执行器，网页发单、执行器干活。你要做的只有两件事——网页配账密、电脑装一次机。"
        "装完之后这台电脑就不用再碰了。")

# ── P3 全景链路 ──────────────────────────────────────────────────────
P.slide_chain(P.prs, 3, "How it works", "一单成绩导出的完整流转", [
    ("网页发起", "workbench"), ("平台派单", "自动派单"), ("执行器接单", "Windows 服务"),
    ("登 e听说", "托管账密"), ("下载美化", "excel-worker"), ("网页收成品", "已完成历史"),
], hi=2,
    sub_title="你只需要记住",
    subs=[("唯一入口是网页", "配账密 · 发作业 · 收成品"),
          ("电脑只装一次", "装完不再需要人碰"),
          ("账密各家隔离", "每个租户只看自己的")])
P.notes(P.prs.slides[-1],
        "看这条链路：网页发起后平台自动派单，代理商电脑上的执行器领任务，用平台托管的账密登 e听说，"
        "下载美化后传回平台，成品回网页收。橙色高亮就是装在你电脑里的执行器。")

# ── P4 章节 01 ───────────────────────────────────────────────────────
P.slide_section(P.prs, 1, "网页配置 · 3 分钟", [
    "浏览器打开 eai.xiaoxiangzhihui.com 并登录",
    "左侧菜单：系统管理 → 执行器下载",
    "「e听说 凭据」卡：填账密 → 保存",
    "点「下载安装包」（约 250 MB）",
], total=4, domain="eai.xiaoxiangzhihui.com")
P.notes(P.prs.slides[-1], "第一章，网页配置，3 分钟，四步走完。")

# ── P5 登录 ──────────────────────────────────────────────────────────
media(5, "Step 1 · Login", "第 1 步 · 打开平台并登录", "ann-login.png", [
    ("输入账号", "平台方发放的运营账号（演示号 zhang.ops），输入工号或邮箱。"),
    ("输入密码", "初始密码找平台方要；输错会提示，可重试。"),
    ("点「登 录」", "登录成功进入「今日待办」页。"),
], caption="截图 2026-09-08 实拍 · eai.xiaoxiangzhihui.com",
    footnote="登录后第一页是「今日待办」；本轮成绩主路径是按日期导出并美化，不发家长群。")
P.notes(P.prs.slides[-1],
        "新电脑浏览器打开网址，输入平台方发放的账号密码，点登录。进来第一页是今日待办，有任务会列在这里。")

# ── P6 进装机页 ──────────────────────────────────────────────────────
media(6, "Step 2 · Open page", "第 2 步 · 系统管理 → 执行器下载", "ann-menu.png", [
    ("点开「系统管理」", "左侧菜单栏下部，点一下展开子菜单。"),
    ("点「执行器下载」", "子菜单第三项，点开就是装机页。"),
], caption="截图 2026-09-08 实拍 · 左侧导航",
    footnote="截图时页面还停在「今日待办」；点开执行器下载后，右侧主区切换为装机页。")
P.notes(P.prs.slides[-1],
        "在左侧菜单点开系统管理，再点执行器下载。这一页就是装机页，后面配账密、下安装包都在这里。")

# ── P7 凭据卡 + 下载 ─────────────────────────────────────────────────
media(7, "Step 3 · Credentials & download", "第 3 步 · 配置 e听说 凭据，点下载安装包", "ann-executor.png", [
    ("看「已配置」徽标", "卡片右上角显示已配置 = 平台已存好这家代理商的账密。"),
    ("填 / 改账密", "输入该代理商的 e听说 账号密码，可下拉选历史账号。"),
    ("点「保存」", "随时生效；以后 e听说 改密码也回这里改。"),
    ("点「下载安装包」", "约 250 MB，下载几分钟；安装包零依赖，装完即用。"),
], caption="截图 2026-09-08 实拍 · 执行器下载页 · 账号已脱敏",
    footnote="这步是告诉平台「这家代理商用什么账密登 e听说」；安装指引文件也在下载区。")
P.notes(P.prs.slides[-1],
        "右上角已配置徽标亮着，说明平台已经存好了账密。要填或要改，在卡里输入账号密码点保存，随时生效。"
        "然后点下载安装包，250 MB 要下几分钟，趁这时间正好喝口水。")

# ── P8 章节 02 ───────────────────────────────────────────────────────
P.slide_section(P.prs, 2, "安装 · 3 分钟", [
    "在代理商要装的那台电脑上，双击 EAiScoreWorker-Setup.exe",
    "「执行器配置」页：三项全不用填",
    "点「安装」，等它跑完",
    "自动注册 Windows 服务「小象智汇成绩执行器」",
], total=4, domain="安装包 v0.3.0 · 约 250 MB")
P.notes(P.prs.slides[-1], "第二章，安装，3 分钟。核心就一句：向导里什么都不用填。")

# ── P9 安装向导示意 ──────────────────────────────────────────────────
def slide_wizard(prs, idx):
    s = P.add_slide(prs)
    P.page_chrome(s, idx, "Step 4 · Install wizard")
    P.text(s, Inches(P.MARGIN), Inches(0.95), Inches(12.2), Inches(0.7),
           "安装向导：三项全不用填，直接点安装", 30, P.INK, True).name = "title"
    # 左：向导面板示意
    px, py, pw, ph = 0.55, 1.9, 6.3, 4.55
    P.box(s, Inches(px), Inches(py), Inches(pw), Inches(ph), fill=P.PAPER, line=P.LINE).name = "wizbox"
    P.box(s, Inches(px), Inches(py), Inches(pw), Inches(0.55), fill=P.CREAM).name = "wizbar"
    P.text(s, Inches(px + 0.25), Inches(py + 0.13), Inches(pw - 0.5), Inches(0.3),
           "小象智汇成绩执行器 · 安装向导（示意）", 11, P.MUTED, True)
    fields = [("租户号", "wh-hj（已预填 · 一般不改）"),
              ("e听说 账号", "留空"),
              ("e听说 密码", "留空")]
    for i, (lab, val) in enumerate(fields):
        y = 2.75 + i * 0.78
        P.text(s, Inches(px + 0.3), Inches(y + 0.08), Inches(1.7), Inches(0.35),
               lab, 12, P.INK, True).name = f"wiz{i}lab"
        P.box(s, Inches(px + 2.1), Inches(y), Inches(3.4), Inches(0.5),
              fill=P.PAPER, line=P.LINE).name = f"wiz{i}box"
        P.text(s, Inches(px + 2.28), Inches(y + 0.11), Inches(3.1), Inches(0.3),
               val, 11, P.MUTED).name = f"wiz{i}val"
    P.text(s, Inches(px + 0.3), Inches(5.35), Inches(3.4), Inches(0.45),
           "三项全不用填", 17, P.CORAL_DEEP, True).name = "wizstamp"
    btn = P.box(s, Inches(px + 4.2), Inches(5.62), Inches(1.55), Inches(0.52), fill=P.CORAL)
    btn.name = "wizbtn"
    P.text(s, Inches(px + 4.2), Inches(5.71), Inches(1.55), Inches(0.35),
           "安 装", 13, P.PAPER, True, P.PP_ALIGN.CENTER).name = "wizbtn_t"
    # 右：三条说明
    side = [
        ("租户号已预填", "按当前登录租户自动带出（如 wh-hj），一般不改。"),
        ("账密留空", "平台「e听说 凭据」卡已配置，执行器自动拉取托管账密。"),
        ("点「安装」", "平台地址与密钥已内置；装完自动注册 Windows 服务。"),
    ]
    for i, (t, d) in enumerate(side):
        y = 2.0 + i * 1.08
        P.box(s, Inches(7.35), Inches(y + 0.08), Inches(0.14), Inches(0.14), fill=P.CORAL).name = f"wz{i}dot"
        P.text(s, Inches(7.67), Inches(y), Inches(4.6), Inches(0.35), t, 14, P.INK, True).name = f"wz{i}h"
        P.check_fit(d, 10.5, 4.6, 0.62, label=f"wz{i}d")
        P.text(s, Inches(7.67), Inches(y + 0.38), Inches(4.6), Inches(0.62),
               d, 10.5, P.INK_SOFT, spacing=1.3).name = f"wz{i}d"
    P.text(s, Inches(P.MARGIN), Inches(6.55), Inches(12.2), Inches(0.3),
           "图示为讲解用重绘的向导示意，非真实截图；兜底填法见网页「安装指引」。",
           9, P.MUTED, font=P.FONT_MONO).name = "wizfn"
    return s

slide_wizard(P.prs, 9)
P.notes(P.prs.slides[-1],
        "双击安装包一路下一步，走到执行器配置页：租户号已预填，账号密码平台托管、全部留空，直接点安装。"
        "平台地址和密钥已经内置，不需要填。")

# ── P10 装完即自动运行 ───────────────────────────────────────────────
P.slide_cards(P.prs, 10, "Auto-run", "装完即自动运行，这台电脑不用再管", [
    ("开机自启", "注册为 Windows 服务\n「小象智汇成绩执行器」\n电脑重启它自己起来。"),
    ("崩溃自动重启", "服务异常自动恢复\n不需要人工干预。"),
    ("登录态自动维护", "e听说 会话过期\n自动用托管账密重登。"),
    ("桌面零图标", "没有任何图标要点\n日常不占用这台电脑。"),
], cols=4, card_h=3.3)
P.notes(P.prs.slides[-1],
        "装完它就是一个 Windows 服务：开机自启、崩溃自动重启、登录态过期自动重登。桌面没有任何图标要点，日常不用管。")

# ── P11 章节 03 ──────────────────────────────────────────────────────
P.slide_section(P.prs, 3, "验证 · 2 分钟", [
    "开始菜单 →「小象智汇成绩执行器」→「查看日志」",
    "应看到：拉到托管账密 → 登录 e听说 成功 → 待命",
    "无红色报错、无连续 401 / 403 = 装机完成",
], total=4, domain="开始菜单搜索「小象智汇」可直达")
P.notes(P.prs.slides[-1], "第三章，验证，2 分钟。装得好不好，日志说了算。")

# ── P12 日志判定 ─────────────────────────────────────────────────────
def slide_logcheck(prs, idx):
    s = P.add_slide(prs)
    P.page_chrome(s, idx, "Step 6 · Verify log")
    P.text(s, Inches(P.MARGIN), Inches(0.95), Inches(12.2), Inches(0.7),
           "验证：三行日志都在，就是装好了", 30, P.INK, True).name = "title"
    # 左：日志面板示意
    px, py, pw, ph = 0.55, 1.95, 6.6, 2.95
    P.box(s, Inches(px), Inches(py), Inches(pw), Inches(ph), fill=P.INK).name = "logbox"
    P.text(s, Inches(px + 0.3), Inches(py + 0.18), Inches(pw - 0.6), Inches(0.3),
           "服务日志（示意）", 9, P.CREAM, font=P.FONT_MONO).name = "logcap"
    logs = [
        "[INFO] 从平台拉取托管账密 ......... 成功",
        "[INFO] 登录 e听说 ................. 成功",
        "[INFO] 执行器待命，等待平台任务",
    ]
    P.text(s, Inches(px + 0.3), Inches(py + 0.62), Inches(pw - 0.6), Inches(2.0),
           logs, 11.5, P.PAPER, spacing=1.6, font=P.FONT_MONO).name = "loglines"
    P.box(s, Inches(px), Inches(5.2), Inches(pw), Inches(0.85), fill=P.CREAM).name = "logwarn"
    P.text(s, Inches(px + 0.3), Inches(5.34), Inches(pw - 0.6), Inches(0.6),
           "红色报错、连续 401 / 403 = 没装好：截图发给平台方", 13, P.INK, True,
           spacing=1.25).name = "logwarn_t"
    # 右：三条判定
    side = [
        ("在哪看", "开始菜单 →「小象智汇成绩执行器」→「查看日志」。"),
        ("什么算就绪", "依次看到：拉到托管账密 → 登录成功 → 待命。"),
        ("什么算有问题", "红色报错或连续 401 / 403：截图发平台方，不要自己重装。"),
    ]
    for i, (t, d) in enumerate(side):
        y = 2.0 + i * 1.08
        P.box(s, Inches(7.55), Inches(y + 0.08), Inches(0.14), Inches(0.14), fill=P.CORAL).name = f"lg{i}dot"
        P.text(s, Inches(7.87), Inches(y), Inches(4.4), Inches(0.35), t, 14, P.INK, True).name = f"lg{i}h"
        P.check_fit(d, 10.5, 4.4, 0.62, label=f"lg{i}d")
        P.text(s, Inches(7.87), Inches(y + 0.38), Inches(4.4), Inches(0.62),
               d, 10.5, P.INK_SOFT, spacing=1.3).name = f"lg{i}d"
    P.text(s, Inches(P.MARGIN), Inches(6.55), Inches(12.2), Inches(0.3),
           "日志为示意重绘，判定口径与安装包 v0.3.0 一致；更多排查见网页「安装指引」。",
           9, P.MUTED, font=P.FONT_MONO).name = "logfn"
    return s

slide_logcheck(P.prs, 12)
P.notes(P.prs.slides[-1],
        "开始菜单找小象智汇成绩执行器，点查看日志。依次看到拉取账密、登录成功、待命，三行都在、没有红色报错，就是装好了。"
        "有红色报错或连续 401、403，截图找平台方，别自己重装。")

# ── P13 章节 04 ──────────────────────────────────────────────────────
P.slide_section(P.prs, 4, "日常使用 · 全自动", [
    "网页发起一次成绩导出，执行器自动接单",
    "「已完成历史」里点详情，zip 下载全部成品",
    "e听说 改密码：只改网页凭据卡，电脑零操作",
], total=4, domain="从这里开始，装机电脑不再需要人碰")
P.notes(P.prs.slides[-1], "第四章，日常使用。从这里开始，这台电脑不再需要人碰，所有动作都在网页。")

# ── P14 发起导出 ─────────────────────────────────────────────────────
media(14, "Daily · Launch", "日常 · 网页发起一次成绩导出", "ann-create.png", [
    ("选「要做什么」", "成绩智能体 · 下载并处理成绩（默认已选好）。"),
    ("选学校", "输入校名检索，如：巴东县水布垭中学。"),
    ("选统计日期", "要哪段成绩就选哪段，默认最近一周。"),
    ("点「发起」", "执行器自动接单开跑；班级留空 = 该校全部学生班。"),
], caption="截图 2026-09-08 实拍 · 智能体作业 → 发起作业",
    footnote="左上「批量发起」可一次发多个班级；发起后列表里状态走到「导出中」。")
P.notes(P.prs.slides[-1],
        "日常唯一动作：智能体作业页点发起作业。选成绩智能体、检索学校、选统计日期，班级留空就是全校，点发起。"
        "执行器马上自动接单，列表状态变成导出中。")

# ── P15 收成品 ───────────────────────────────────────────────────────
media(15, "Daily · Collect", "日常 · 几分钟后，网页收成品", "ann-done.png", [
    ("看「执行状态」", "完全成功（绿标）= 成品就绪；失败进「待人工干预」。"),
    ("点「详情」", "zip 打包下载这次导出的全部结果文件。"),
], caption="截图 2026-09-08 实拍 · 已完成历史（本代理商）",
    footnote="单条与批量终态都在本页；批量行用「查看子作业」进子列表。")
P.notes(P.prs.slides[-1],
        "几分钟后到已完成历史看结果：状态完全成功，点详情就能 zip 下载全部成品文件。"
        "失败的任务会出现在待人工干预，等人工处理。")

# ── P16 两个日常场景 ─────────────────────────────────────────────────
P.slide_versus(P.prs, 16, "Daily scenarios", "两个常见场景，都不用碰装机电脑",
    ("e听说 改了密码", [
        "网页「执行器下载」页，找到 e听说 凭据卡",
        "填新账号密码，点保存，随时生效",
        "下个任务自动用新密码重登",
        "装机电脑零操作",
    ]),
    ("新增一家代理商", [
        "新租户在网页配自己的 e听说 账密",
        "新电脑按本教材第 1-3 章装一次执行器",
        "各家账密互相看不见",
    ]))
P.notes(P.prs.slides[-1],
        "两个常见场景：e听说改密码，网页卡里改完保存就行，电脑零操作；新增代理商，新租户配账密、新电脑装一次机，各家账密互相看不见。")

# ── P17 尾页三句话 ───────────────────────────────────────────────────
def slide_takeaway(prs):
    s = P.add_slide(prs)
    P.box(s, 0, 0, P.SW, P.SH, fill=P.INK)
    P.box(s, Inches(0.9), Inches(1.45), Inches(0.35), Emu(19050), fill=P.CORAL).name = "cdivider"
    P.text(s, Inches(0.9), Inches(1.0), Inches(8), Inches(0.4),
           "TAKEAWAY · 三句话带走", 13, P.CORAL, True).name = "kicker"
    lines = [
        ("网页是唯一入口", "eai.xiaoxiangzhihui.com：配 e听说 账密 · 发作业 · 收成品"),
        ("装机只做一次", "下载安装包，一路下一步，向导里什么都不用填"),
        ("之后全自动", "改密码只改网页凭据卡，装机电脑零操作"),
    ]
    for i, (h, d) in enumerate(lines):
        y = 2.15 + i * 1.25
        P.text(s, Inches(0.9), Inches(y), Inches(1.0), Inches(0.7),
               str(i + 1), 30, P.CORAL, True, font=P.FONT_MONO).name = f"tk{i}no"
        P.text(s, Inches(1.85), Inches(y), Inches(10.4), Inches(0.45),
               h, 21, P.PAPER, True).name = f"tk{i}h"
        P.text(s, Inches(1.85), Inches(y + 0.52), Inches(10.4), Inches(0.4),
               d, 13, P.CREAM).name = f"tk{i}d"
    P.text(s, Inches(0.9), Inches(6.35), Inches(11), Inches(0.4),
           "装一次，之后全自动。", 20, P.CORAL, True).name = "slogan"
    P.text(s, Inches(0.9), Inches(6.95), Inches(11), Inches(0.3),
           "eai.xiaoxiangzhihui.com  ·  小象智汇平台方  ·  2026-09",
           10, P.CREAM, font=P.FONT_MONO).name = "cfoot"
    return s

def Emu_bar():
    from pptx.util import Emu
    return Emu(19050)

slide_takeaway(P.prs)
P.notes(P.prs.slides[-1],
        "收尾三句话：网页是唯一入口；装机只做一次，什么都不用填；之后全自动，改密码只改网页。有问题随时找平台方。")

# ── 转场 + 保存 ──────────────────────────────────────────────────────
P.set_transition(P.prs, "fade")
P.prs.save(OUT)
print("OK", OUT, len(P.prs.slides._sldIdLst), "slides")
