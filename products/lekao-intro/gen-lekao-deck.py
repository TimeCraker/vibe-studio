# -*- coding: utf-8 -*-
"""gen-lekao-deck.py - LeKao product intro deck (11 pages) + narration script.json.

Facts sourced ONLY from lekao/README.md. Brand palette extracted from
lekao/src/app/globals.css (see primitives.py header in this directory).
Narration segments live in NARRATION below; the same data feeds speaker
notes (per page) and script.json, so deck and script stay in sync.

Usage: python gen-lekao-deck.py
Outputs: deck.pptx, script.json (same directory as this script).
Console output is English only (GBK console).
"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx.util import Inches

import primitives as P

OUT_DIR = os.path.dirname(os.path.abspath(__file__))

# ── narration: page -> [segment texts] (<=20 chars each after stripping) ──
NARRATION = {
    1: ["LeKao，AI 智能助教助手。", "K12 机构助教的 AI 工作台。"],
    2: ["每次课，抄错题、写小结、写反馈。", "三件事，吃掉 1 到 2 个小时。"],
    3: ["这三件事，LeKao 一次接走。", "小结、反馈、错题集，全是成品。"],
    4: ["讲义拍照上传，AI 先读图。", "学科、章节、知识点，自动认。", "10 条反馈加一段小结，直接出。"],
    5: ["作业堂测，同样传图生成。", "逐题反馈，再加总体小结。", "字数四挡，标准到三倍随意选。"],
    6: ["全班错题集，五分钟成档。", "Excel 模板填好，图片拖进来。", "浏览器里直接生成 Word。"],
    7: ["T-Coin 哪里来？每天签到。", "连续签到，从 2 币涨到 8 币。", "生成一次花 1 币，失败自动退回。"],
    8: ["手机上传，图片自动压缩。", "生成流式显示，边跑边出。", "历史存在本地，随时回看。"],
    9: ["邮箱加验证码，注册就能用。", "传讲义，选功能，拿结果。"],
    10: ["记住一句话就够了。", "把每次课的 1 到 2 小时，还给教学。"],
    11: ["打开官网，现在就能试。", "每次课，省回 1 到 2 小时。"],
}

BADGE = "LEKAO"


def bottom_band(slide, y, h, lines, big=None):
    """Page-bottom takeaway band: fills the lower whitespace of sparse pages."""
    P.box(slide, Inches(0.55), Inches(y), Inches(12.23), Inches(h), fill=P.CREAM)
    if big:
        P.text(slide, Inches(0.9), Inches(y), Inches(2.7), Inches(h),
               big, 40, P.CORAL_DEEP, True, anchor=P.MSO_ANCHOR.MIDDLE)
        P.text(slide, Inches(3.8), Inches(y), Inches(8.7), Inches(h),
               lines, 13, P.INK_SOFT, spacing=1.4, anchor=P.MSO_ANCHOR.MIDDLE)
    else:
        P.text(slide, Inches(0.9), Inches(y), Inches(11.5), Inches(h),
               lines, 13, P.INK, align=P.PP_ALIGN.CENTER, anchor=P.MSO_ANCHOR.MIDDLE)


def main():
    prs = P.new_deck()
    slides = []

    # P1 cover
    meta = {
        "kicker": "LEKAO PRODUCT INTRO",
        "title": "LeKao",
        "subtitle": "AI 智能助教助手 · K12 机构助教工作台",
        "tagline": "把抄错题、写小结、写反馈交给 AI，每次课省 1-2 小时",
        "domain": "lekao.asterforge.top",
        "version": "",
        "date": "2026-08",
        "author": "ASTERFORGE",
    }
    slides.append(P.slide_cover(prs, meta))

    # P2 pain point
    s2 = P.slide_rows(
        prs, 2, "PAIN POINT", "每次课，1-2 小时耗在重复劳动上", [
            ("抄错题", "每节课后，把学生错题重新整理一遍"),
            ("写小结", "每节课都要写一段课堂小结"),
            ("写反馈", "作业要逐题给出反馈"),
        ],
    )
    bottom_band(
        s2, 5.05, 1.65,
        [("这是助教每上一节课，都要在三个环节重复花掉的时间。", 13, P.INK_SOFT, False),
         ("LeKao 要接走的，正是这三件事。", 13, P.INK, True)],
        big="1-2 小时",
    )
    slides.append(s2)

    # P3 what it is
    s3 = P.slide_cards(
        prs, 3, "PRODUCT", "K12 机构助教的 AI 工作台", [
            ("课堂小结", "上传讲义图片\nAI 识别学科、章节、知识点\n生成 10 条反馈 + 一段小结"),
            ("作业 / 堂测反馈", "同样上传图片生成\n逐题反馈 + 总体小结\n字数 4 挡可选"),
            ("错题集 Word", "Excel 模板 + 题目图片\n5 分钟生成全班错题文档\n浏览器内直接完成"),
        ], cols=3, card_h=2.9,
    )
    bottom_band(s3, 5.85, 0.85, ["三个功能共用一个工作台，浏览器打开即可使用。"])
    slides.append(s3)

    # P4 feature 1: class summary
    slides.append(P.slide_chain(
        prs, 4, "FEATURE 01 · CLASS SUMMARY", "课堂小结：一张讲义图，出一份成品",
        nodes=[
            ("上传讲义图片", "拍照 / 截图均可"),
            ("AI 读图识别", "学科 · 章节 · 知识点"),
            ("AI 生成", "10 条反馈 + 一段小结"),
            ("直接能用", "输出格式固定"),
        ], hi=2,
        sub_title="为什么每次都规整",
        subs=[
            ("课堂情况逐条列明", "10 条反馈"),
            ("整堂课一段话收束", "一段小结"),
            ("输出结构每次一致", "格式硬约束"),
        ],
    ))

    # P5 feature 2: homework feedback
    s5 = P.slide_cards(
        prs, 5, "FEATURE 02 · FEEDBACK", "作业与堂测：逐题反馈，字数你定", [
            ("逐题反馈", "上传讲义 / 试卷图片\nAI 逐题生成讲评\n反馈落到每一小题"),
            ("总体小结", "作业整体情况\n自动收成一段总结\n省掉自己归纳"),
            ("4 挡字数", "标准 / 1.5 倍 / 2 倍 / 3 倍\n同一份反馈\n长度随场景伸缩"),
        ], cols=3, card_h=2.9,
    )
    bottom_band(s5, 5.85, 0.85, ["一次生成两样：逐题讲评 + 总体小结，复制即可用。"])
    slides.append(s5)

    # P6 feature 3: mistake book
    slides.append(P.slide_chain(
        prs, 6, "FEATURE 03 · MISTAKE BOOK", "错题集：五分钟，全班一本 Word",
        nodes=[
            ("填 Excel 模板", "按模板登记错题"),
            ("拖入题目图片", "批量拖拽上传"),
            ("浏览器内生成", "5 分钟全班一份"),
            ("导出 Word", "直接打印分发"),
        ], hi=2,
        sub_title="这一步在本地浏览器完成",
        subs=[
            ("docx 在浏览器本地拼装", "纯前端生成"),
            ("错题集全程不扣币", "不消耗 T-Coin"),
            ("打印、发群拿来即用", "全班一本"),
        ],
    ))

    # P7 T-Coin
    slides.append(P.slide_numbers(
        prs, 7, "T-COIN", "T-Coin：签到领币，生成花币", [
            ("5 币", "注册即送", "新用户注册自动补发 5 枚"),
            ("2-8 币", "每日签到可得", "连续签到 7 日，奖励逐日递增"),
            ("1 币", "AI 生成一次", "课堂小结与作业反馈均为 1 币"),
            ("自动退回", "生成失败不白耗", "失败那次消耗的币自动退回"),
        ], cols=2,
    ))

    # P8 experience
    slides.append(P.slide_rows(
        prs, 8, "EXPERIENCE", "手机上一样顺手", [
            ("图片自动压缩", "移动端自动降级：1280px / 0.4MB 以内，大图上传不卡顿"),
            ("流式输出", "生成内容实时逐段显示，边生成边看"),
            ("本地历史", "生成记录存在浏览器 IndexedDB，随时回看"),
            ("登录方式", "邮箱密码注册 + 邮箱验证码，换设备照常登录"),
        ],
    ))

    # P9 getting started
    slides.append(P.slide_chain(
        prs, 9, "GETTING STARTED", "三步上手",
        nodes=[
            ("注册登录", "邮箱 + 验证码"),
            ("上传讲义图", "小结 / 反馈 / 错题集"),
            ("拿结果", "Word 成品直接用"),
        ], hi=1,
        sub_title="现在就能体验",
        subs=[
            ("lekao.asterforge.top", "在线体验"),
            ("登录后可查站内使用指南", "使用指南"),
        ],
    ))

    # P10 takeaway quote
    slides.append(P.slide_quote(
        prs,
        "把每次课的 1-2 小时，还给教学本身。",
        "LeKao · lekao.asterforge.top",
    ))

    # P11 closing
    s = P.slide_closing(
        prs,
        {"domain": "lekao.asterforge.top"},
        ["lekao.asterforge.top", "LeKao · AI 智能助教助手", "© 2026 ASTERFORGE"],
        slogan="每次课，省回 1-2 小时。",
    )
    for sp in s.shapes:  # brand mark: TC -> LK
        if sp.name == "logot":
            sp.text_frame.paragraphs[0].runs[0].text = "LK"
    slides.append(s)

    # speaker notes = page narration
    for i, slide in enumerate(slides, 1):
        P.notes(slide, "\n".join(NARRATION[i]))

    deck_path = os.path.join(OUT_DIR, "deck.pptx")
    prs.save(deck_path)
    print(f"OK deck: {deck_path} pages={len(slides)}")

    # script.json from the same narration data
    segments = []
    sid = 1
    for page in sorted(NARRATION):
        for text in NARRATION[page]:
            segments.append({"id": sid, "kind": "subtitle", "text": text,
                             "ref": f"page {page}"})
            sid += 1
    script = {
        "meta": {
            "title": "LeKao 产品介绍视频",
            "lang": "zh",
            "budgetSeconds": 75,
            "charsPerSecond": 4.5,
            "bannedWords": ["赋能", "闭环", "沉淀", "矩阵", "生态", "一站式",
                            "端到端", "付费", "充值", "价格", "收费", "免费",
                            "退款", "成本", "花钱"],
        },
        "segments": segments,
    }
    script_path = os.path.join(OUT_DIR, "script.json")
    with open(script_path, "w", encoding="utf-8") as f:
        json.dump(script, f, ensure_ascii=False, indent=2)
    print(f"OK script: {script_path} segments={len(segments)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
